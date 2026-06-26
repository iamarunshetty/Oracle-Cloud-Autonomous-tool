"""
Generate sample Excel templates for bulk operations.

Run directly to regenerate templates:
    python -m app.services.template_generator
"""
from __future__ import annotations

from pathlib import Path


def create_bulk_user_template(output_path: str | Path | None = None) -> Path:
    """
    Write a bulk-user-upload Excel template.
    Returns the path of the written file.
    """
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill
    except ImportError as exc:
        raise RuntimeError("openpyxl is required: pip install openpyxl") from exc

    if output_path is None:
        output_path = Path(__file__).resolve().parent.parent.parent / "templates" / "bulk_user_upload_template.xlsx"

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Users"

    headers = [
        "username",
        "first_name",
        "last_name",
        "email",
        "department",
        "job_title",
        "manager_username",
    ]
    required = {"username", "first_name", "last_name", "email"}

    ws.append(headers)
    for cell in ws[1]:
        if cell.value in required:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="003366")
        else:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="B8CCE4")

    # Sample data rows
    samples = [
        ["jdoe", "John", "Doe", "jdoe@example.com", "Finance", "Analyst", ""],
        ["asmith", "Alice", "Smith", "asmith@example.com", "IT", "Engineer", "jdoe"],
    ]
    for row in samples:
        ws.append(row)

    # Column widths
    widths = [18, 15, 15, 28, 18, 20, 20]
    for col, width in zip(ws.iter_cols(min_col=1, max_col=len(headers)), widths):
        ws.column_dimensions[col[0].column_letter].width = width

    # Instructions sheet
    ws2 = wb.create_sheet("Instructions")
    ws2.append(["Oracle Fusion Bulk User Upload – Instructions"])
    ws2.append([])
    ws2.append(["Required columns (highlighted dark blue): username, first_name, last_name, email"])
    ws2.append(["Optional columns (light blue): department, job_title, manager_username"])
    ws2.append([])
    ws2.append(["Rules:"])
    ws2.append(["- username: must be unique, no spaces"])
    ws2.append(["- email: must be a valid email address"])
    ws2.append(["- manager_username: must match an existing user's username (or leave blank)"])
    ws2.append(["- Do not modify the column headers"])

    wb.save(output_path)
    return output_path


if __name__ == "__main__":
    path = create_bulk_user_template()
    print(f"Template written to: {path}")
